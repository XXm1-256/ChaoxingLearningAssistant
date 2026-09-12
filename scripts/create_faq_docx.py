from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent.parent
SOURCE = ROOT / "常见问题与处理方法.md"
OUTPUT = ROOT / "常见问题与处理方法.docx"
ACCENT = RGBColor.from_string("3D9B67")
DARK = RGBColor.from_string("17392D")
MUTED = RGBColor.from_string("4E6C5C")


def style_run(run, size=11, bold=False, color=DARK):
    run.font.name = "Microsoft YaHei UI"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei UI")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.7)
section.bottom_margin = Inches(0.7)
section.left_margin = Inches(0.78)
section.right_margin = Inches(0.78)

for style_name in ("Normal", "Title", "Heading 1"):
    style = doc.styles[style_name]
    style.font.name = "Microsoft YaHei UI"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei UI")

lines = SOURCE.read_text(encoding="utf-8").splitlines()
for line in lines:
    text = line.strip()
    if not text:
        continue
    if text.startswith("# "):
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.space_after = Pt(8)
        style_run(paragraph.add_run(text[2:]), size=22, bold=True, color=RGBColor(0, 0, 0))
    elif text.startswith("## "):
        if text == "## 左侧没有课程或章节":
            doc.add_page_break()
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(11)
        paragraph.paragraph_format.space_after = Pt(4)
        style_run(paragraph.add_run(text[3:]), size=14, bold=True, color=ACCENT)
    else:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.line_spacing = 1.32
        paragraph.paragraph_format.space_after = Pt(6)
        style_run(paragraph.add_run(text), size=10.7, color=MUTED if text.startswith("适用于") else DARK)

doc.core_properties.title = "学习通课程视频播放助手常见问题与处理方法"
doc.core_properties.subject = "学习通课程视频播放助手 v1.45 常见问题"
doc.core_properties.author = "ChaoxingLearningAssistant"
doc.save(OUTPUT)
print(OUTPUT)
