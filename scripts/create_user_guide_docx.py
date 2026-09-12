from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent.parent
OUTPUT = ROOT / "使用教学.docx"
SCREENSHOT = ROOT / "docs" / "validation" / "v1.37" / "app-main.png"
if not SCREENSHOT.exists():
    SCREENSHOT = ROOT / "docs" / "validation" / "v1.34" / "app-main.png"
ACCENT = "3D9B67"
PALE = "E8F5EC"
DARK = "17392D"
GRID = "D9E4DC"


def set_cell_shading(cell, fill):
    props = cell._tc.get_or_add_tcPr()
    shade = props.find(qn("w:shd"))
    if shade is None:
        shade = OxmlElement("w:shd")
        props.append(shade)
    shade.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=120, start=140, bottom=120, end=140):
    props = cell._tc.get_or_add_tcPr()
    margins = props.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        props.append(margins)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table):
    props = table._tbl.tblPr
    borders = props.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        props.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "6")
        node.set(qn("w:color"), GRID)
        borders.append(node)


def remove_paragraph_border(paragraph):
    props = paragraph._p.get_or_add_pPr()
    border = props.find(qn("w:pBdr"))
    if border is not None:
        props.remove(border)


def set_font(run, name="Microsoft YaHei UI", size=11, bold=False, color="000000"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.35
    if bold_lead and text.startswith(bold_lead):
        set_font(p.add_run(bold_lead), bold=True)
        text = text[len(bold_lead):]
    set_font(p.add_run(text))
    return p


def add_steps(doc, steps):
    table = doc.add_table(rows=len(steps), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(0.62)
    table.columns[1].width = Inches(6.55)
    set_table_borders(table)
    for index, text in enumerate(steps, 1):
        number, detail = table.rows[index - 1].cells
        number.width = Inches(0.62)
        detail.width = Inches(6.55)
        number.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        detail.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_shading(number, ACCENT)
        if index % 2 == 0:
            set_cell_shading(detail, PALE)
        for cell in (number, detail):
            set_cell_margins(cell)
        number.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_font(number.paragraphs[0].add_run(str(index)), size=12, bold=True, color="FFFFFF")
        detail.paragraphs[0].paragraph_format.line_spacing = 1.25
        set_font(detail.paragraphs[0].add_run(text), size=10.8)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.68)
section.bottom_margin = Inches(0.68)
section.left_margin = Inches(0.72)
section.right_margin = Inches(0.72)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Microsoft YaHei UI"
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei UI")
normal.font.size = Pt(11)
for name, size in (("Title", 28), ("Heading 1", 19), ("Heading 2", 14)):
    style = styles[name]
    style.font.name = "Microsoft YaHei UI"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei UI")
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.font.bold = True
    style.paragraph_format.space_before = Pt(16 if name != "Title" else 0)
    style.paragraph_format.space_after = Pt(7)
    style_props = style._element.get_or_add_pPr()
    style_border = style_props.find(qn("w:pBdr"))
    if style_border is not None:
        style_props.remove(style_border)

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.LEFT
remove_paragraph_border(title)
set_font(title.add_run("学习通课程视频播放助手使用教学"), size=28, bold=True)
subtitle = doc.add_paragraph()
subtitle.paragraph_format.space_after = Pt(16)
set_font(subtitle.add_run("适用于 v1.37 Windows 便携版"), size=11.5, color="4E6C5C")

add_body(doc, "这份教学面向第一次使用的同学。程序帮助减少重复点击：识别课程中的真实视频，按同一章节内的视频顺序继续播放，再查找后续未完成视频。课程是否完成仍以学习通页面的实际记录为准。")

doc.add_heading("五步开始使用", level=1)
add_steps(doc, [
    "完整解压下载包，不要直接在压缩包里运行程序。",
    "进入“学习通课程视频播放助手”文件夹，双击 ChaoxingLearningAssistant.exe。",
    "在程序中的学习通页面完成登录，并打开需要学习的课程。",
    "点顶部的“开始”。功能开启后会出现提示，右侧会显示当前视频和下一视频。",
    "保持程序开启。视频自然结束后，程序会继续寻找下一段真实视频。",
])

if SCREENSHOT.exists():
    doc.add_heading("界面位置", level=1)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(SCREENSHOT), width=Inches(7.0))
    caption = doc.add_paragraph("程序主界面：左侧选择课程与章节，中间显示学习通网页，右侧显示播放状态。")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_after = Pt(10)
    for run in caption.runs:
        set_font(run, size=9.5, color="4E6C5C")

doc.add_page_break()
doc.add_heading("主要按钮", level=1)
buttons = [
    ("开始", "识别当前课程和待播放视频，并尝试开始播放。"),
    ("打开未完成章节", "打开候选章节后再核对真实视频；若视频均已看完、只剩测验或作业，会继续寻找后面的未看完视频。"),
    ("暂停辅助", "保留当前页面，但暂时不再自动查找或切换下一视频。"),
    ("全屏观看", "按 F11 让学习通网页铺满窗口；点顶部“退出全屏”，或按 Esc、F11 返回。"),
    ("刷新网页", "重新载入当前学习通页面。"),
    ("课程首页", "回到学习通课程入口。"),
    ("日志与反馈", "查看程序运行记录。"),
    ("导出反馈包", "保存日志和当时的页面状态，便于反馈问题。"),
]
table = doc.add_table(rows=1, cols=2)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False
table.columns[0].width = Inches(1.25)
table.columns[1].width = Inches(5.55)
set_table_borders(table)
table.rows[0].cells[0].width = Inches(1.25)
table.rows[0].cells[1].width = Inches(5.55)
for cell, label in zip(table.rows[0].cells, ("按钮", "作用")):
    set_cell_shading(cell, DARK)
    set_cell_margins(cell)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_font(cell.paragraphs[0].add_run(label), size=10.5, bold=True, color="FFFFFF")
for index, (button, meaning) in enumerate(buttons):
    row = table.add_row().cells
    row[0].width = Inches(1.25)
    row[1].width = Inches(5.55)
    if index % 2 == 1:
        for cell in row:
            set_cell_shading(cell, PALE)
    for cell in row:
        set_cell_margins(cell)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_font(row[0].paragraphs[0].add_run(button), size=10.5, bold=True)
    set_font(row[1].paragraphs[0].add_run(meaning), size=10.5)

doc.add_heading("播放顺序", level=1)
add_body(doc, "程序只把真实视频计入视频顺序，不会把章节测验、作业、考试、签到或普通资料当成下一视频。同一章节有多个视频时，优先播放同章下一段视频；同章视频全部结束后，才会继续查找后续未完成章节。")
add_body(doc, "章节即使显示未完成，只要其中视频都已看完，也会继续向后寻找，不会停在只剩测验或作业的章节。按钮查找和自然结束后的顺序跳转遵守相同规则。")

doc.add_page_break()
doc.add_heading("遇到问题", level=1)
add_steps(doc, [
    "左侧课程为空时，先进入学习通课程页并点“刷新课程”；页面仍在加载时程序会短暂重试。",
    "先点一次“开始”，再观察右侧的“当前视频”和“下一视频”。",
    "点“打开未完成章节”后留意按钮和顶部提示；如果没有识别到真实视频，先在网页目录打开目标章节再点“开始”。",
    "仍然无法继续时，打开“日志与反馈”查看记录。",
    "点“导出反馈包”，并说明课程页面、发生步骤和当时的视频位置。",
    "公开发送反馈包前，检查并删除课程名称、学校名称、页面地址参数等不便公开的信息。",
])

doc.add_heading("使用边界和免责声明", level=1)
add_body(doc, "本程序只是网页播放辅助工具，不绕过学习通的登录、验证码、人脸识别、风控或学校认证，不伪造学习时长、视频进度、成绩或完成状态，不自动答题，也不代替课程要求的人工操作。遇到验证页面时，程序会暂停并等待同学处理。")
add_body(doc, "本作品与超星学习通官方不存在隶属、授权或合作关系。使用时请遵守学校、课程和平台规则。网页结构、网络状态和平台规则可能变化，程序不能保证学习通最终记录结果；重要课程请在学习通页面人工核对进度。")

doc.core_properties.title = "学习通课程视频播放助手使用教学"
doc.core_properties.subject = "学习通课程视频播放助手 v1.37 使用说明"
doc.core_properties.author = "ChaoxingLearningAssistant"
doc.save(OUTPUT)
print(OUTPUT)
